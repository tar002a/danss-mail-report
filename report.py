"""Daily mail-flow report for diri-alnassir.iq and friendgroup-iq.com.

Pulls 24 hours of Exchange Online message trace through Microsoft Graph, builds
an Excel workbook and emails it with an HTML summary. Runs in GitHub Actions and
signs in with the workflow's OIDC token, so no secret is stored anywhere.

Only metadata is read (sender, recipient, subject, status, size); message bodies
and attachments are never accessed.
"""
import base64
import collections
import datetime as dt
import html
import io
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TENANT_ID = os.environ["TENANT_ID"]
CLIENT_ID = os.environ["CLIENT_ID"]
SENDER = os.environ["REPORT_SENDER"]
RECIPIENTS = [a.strip() for a in os.environ["REPORT_TO"].split(",") if a.strip()]
INTERNAL = {d.strip().lower() for d in os.environ.get(
    "INTERNAL_DOMAINS", "diri-alnassir.iq,friendgroup-iq.com").split(",") if d.strip()}

BAGHDAD = dt.timezone(dt.timedelta(hours=3))
GRAPH = "https://graph.microsoft.com"
MAX_ATTACHMENT = 3 * 1024 * 1024  # sendMail limit for inline attachments

FREEMAIL = {
    "gmail.com", "googlemail.com", "yahoo.com", "yahoo.co.uk", "ymail.com", "hotmail.com",
    "hotmail.co.uk", "outlook.com", "live.com", "msn.com", "icloud.com", "me.com", "aol.com",
    "mail.com", "mail.ru", "yandex.com", "yandex.ru", "proton.me", "protonmail.com", "gmx.com",
    "163.com", "126.com", "qq.com", "sina.com", "zoho.com",
}
# Well-known names a typo-squatted domain is compared against
LOOKALIKE_REFS = INTERNAL | {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "live.com", "icloud.com"}

STATUS_LABEL = {
    "delivered": "Delivered", "failed": "Failed", "pending": "Pending", "expanded": "Expanded",
    "quarantined": "Quarantined", "filteredAsSpam": "Filtered as spam", "gettingStatus": "Getting status",
}
BAD_STATUS = {"failed", "quarantined", "filteredAsSpam"}


# --- auth and Graph -------------------------------------------------------

def _json(req):
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.load(r)


def github_oidc_token():
    url = os.environ["ACTIONS_ID_TOKEN_REQUEST_URL"] + "&audience=api://AzureADTokenExchange"
    req = urllib.request.Request(url, headers={
        "Authorization": "Bearer " + os.environ["ACTIONS_ID_TOKEN_REQUEST_TOKEN"]})
    return _json(req)["value"]


def graph_token():
    assertion = github_oidc_token()
    body = urllib.parse.urlencode({
        "client_id": CLIENT_ID,
        "scope": GRAPH + "/.default",
        "grant_type": "client_credentials",
        "client_assertion_type": "urn:ietf:params:oauth:client-assertion-type:jwt-bearer",
        "client_assertion": assertion,
    }).encode()
    try:
        return _json(urllib.request.Request(
            f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token", body))["access_token"]
    except urllib.error.HTTPError as e:
        payload = assertion.split(".")[1]
        claims = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
        print("Token request failed:", e.read().decode(errors="replace"), file=sys.stderr)
        print("OIDC subject presented:", claims.get("sub"), file=sys.stderr)
        raise


def graph(token, url, data=None):
    headers = {"Authorization": "Bearer " + token}
    if data is not None:
        data = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    for attempt in range(6):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, data, headers), timeout=180) as r:
                raw = r.read()
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as e:
            if e.code in (429, 500, 502, 503, 504) and attempt < 5:
                time.sleep(int(e.headers.get("Retry-After") or 30))
                continue
            print(f"Graph {e.code} for {url.split('?')[0]}:", e.read().decode(errors="replace"), file=sys.stderr)
            raise


def message_traces(token, start, end):
    # The API accepts at most 10 days per query; one day is well inside that.
    flt = f"receivedDateTime ge {_z(start)} and receivedDateTime le {_z(end)}"
    url = f"{GRAPH}/beta/admin/exchange/tracing/messageTraces?" + urllib.parse.urlencode(
        {"$filter": flt, "$top": "5000"}, quote_via=urllib.parse.quote)
    while url:
        page = graph(token, url)
        yield from page.get("value", [])
        url = page.get("@odata.nextLink")


def _z(t):
    return t.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# --- classification -------------------------------------------------------

def domain_of(addr):
    return addr.rsplit("@", 1)[-1].lower() if "@" in addr else ""


def is_internal(addr):
    return domain_of(addr) in INTERNAL


def _distance(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def lookalike_of(domain):
    """Return the well-known domain this one imitates, if any."""
    if not domain or domain in INTERNAL or domain in FREEMAIL:
        return None
    label = domain.split(".")[0]
    for ref in LOOKALIKE_REFS:
        if _distance(domain, ref) <= 2:
            return ref
        ref_label = ref.split(".")[0]
        variants = {ref_label}
        if ref in INTERNAL:  # friendgroup-iq -> friendgroupiq, friendgroup
            variants |= {ref_label.replace("-", ""), ref_label.split("-")[0]}
        for v in variants:
            if len(v) < 5:
                continue
            if label != ref_label and _distance(label, v) <= 2:
                return ref
            if label == v and ref in INTERNAL:  # same name, different TLD
                return ref
    return None


def classify(t):
    sender = (t.get("senderAddress") or "").lower()
    rcpt = (t.get("recipientAddress") or "").lower()
    s_int, r_int = is_internal(sender), is_internal(rcpt)
    direction = ("Internal" if s_int and r_int else "Outbound" if s_int
                 else "Inbound" if r_int else "External relay")
    status = t.get("status") or ""
    flags = []
    if direction == "Outbound":
        d = domain_of(rcpt)
        if d in FREEMAIL:
            flags.append("Personal webmail")
        ref = lookalike_of(d)
        if ref:
            flags.append(f"Look-alike of {ref}")
        if status in BAD_STATUS:
            flags.append("Outbound not delivered")
    elif direction == "Internal" and status in BAD_STATUS:
        flags.append("Internal not delivered")
    elif direction == "External relay":
        flags.append("Relayed via tenant")
    received = dt.datetime.fromisoformat(t["receivedDateTime"].replace("Z", "+00:00")).astimezone(BAGHDAD)
    return {
        "time": received, "direction": direction, "sender": sender or "(system)", "recipient": rcpt,
        "subject": t.get("subject") or "", "status": STATUS_LABEL.get(status, status),
        "size_kb": round((t.get("size") or 0) / 1024, 1), "flags": "; ".join(flags),
        "message_id": t.get("messageId") or "",
    }


def mailbox_stats(rows):
    stats = collections.defaultdict(lambda: collections.Counter())
    sent_ids = collections.defaultdict(set)
    for r in rows:
        if is_internal(r["sender"]):
            sent_ids[r["sender"]].add(r["message_id"])
            if r["direction"] == "Outbound":
                stats[r["sender"]]["ext_rcpts"] += 1
                stats[r["sender"]]["ext_domains_" + domain_of(r["recipient"])] += 1
            if r["flags"]:
                stats[r["sender"]]["flagged"] += 1
        if is_internal(r["recipient"]):
            stats[r["recipient"]]["recv_ext" if r["direction"] == "Inbound" else "recv_int"] += 1
    out = []
    for addr in set(stats) | set(sent_ids):
        s = stats[addr]
        out.append({
            "mailbox": addr, "sent": len(sent_ids[addr]), "ext_rcpts": s["ext_rcpts"],
            "ext_domains": len([k for k in s if k.startswith("ext_domains_")]),
            "recv_int": s["recv_int"], "recv_ext": s["recv_ext"], "flagged": s["flagged"],
        })
    out.sort(key=lambda m: (-(m["sent"] + m["recv_int"] + m["recv_ext"]), m["mailbox"]))
    return out


# --- outputs --------------------------------------------------------------

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")


def _sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def build_workbook(rows, boxes, start, end, kpi):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Daily mail-flow report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{start:%d %b %Y %H:%M} – {end:%d %b %Y %H:%M} (Baghdad time)"])
    ws.append([])
    for k, v in kpi.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 14

    _sheet(wb.create_sheet("Mailboxes"),
           ["Mailbox", "Messages sent", "External recipients", "External domains",
            "Received internal", "Received external", "Flagged"],
           [[b["mailbox"], b["sent"], b["ext_rcpts"], b["ext_domains"], b["recv_int"],
             b["recv_ext"], b["flagged"]] for b in boxes],
           [40, 14, 18, 16, 16, 16, 10])

    cols = ["Time (Baghdad)", "Direction", "Sender", "Recipient", "Subject", "Status", "Size KB", "Flags"]
    widths = [17, 13, 36, 36, 60, 15, 9, 30]

    def as_row(r):
        return [r["time"].replace(tzinfo=None), r["direction"], r["sender"], r["recipient"],
                r["subject"], r["status"], r["size_kb"], r["flags"]]

    flagged = [r for r in rows if r["flags"]]
    for name, subset in (("Flags", flagged), ("All messages", rows)):
        sh = wb.create_sheet(name)
        _sheet(sh, cols, [as_row(r) for r in subset], widths)
        for (cell,) in sh.iter_rows(min_row=2, max_col=1):
            cell.number_format = "yyyy-mm-dd hh:mm"
        for row in sh.iter_rows(min_row=2):
            row[4].alignment = Alignment(horizontal="general", readingOrder=0)
            if row[7].value:
                for c in row:
                    c.fill = FLAG_FILL
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _table(headers, rows):
    th = "".join(f'<th style="background:#1F3864;color:#fff;padding:4px 8px;text-align:left">{h}</th>'
                 for h in headers)
    trs = "".join("<tr>" + "".join(
        f'<td dir="auto" style="border-bottom:1px solid #ddd;padding:4px 8px">{html.escape(str(c))}</td>'
        for c in r) + "</tr>" for r in rows)
    return f'<table style="border-collapse:collapse;font-size:13px">{th and "<tr>" + th + "</tr>"}{trs}</table>'


def build_html(rows, boxes, start, end, kpi):
    flagged = [r for r in rows if r["flags"]]
    parts = [
        '<div style="font-family:Segoe UI,Arial,sans-serif;color:#222">',
        "<h2 style=\"margin:0 0 4px\">Daily mail-flow report</h2>",
        f"<p style=\"margin:0 0 12px;color:#555\">{start:%d %b %Y %H:%M} – {end:%d %b %Y %H:%M} "
        "Baghdad time · diri-alnassir.iq and friendgroup-iq.com</p>",
        _table(["Measure", "Count"], kpi.items()),
        f"<h3>Flagged ({len(flagged)})</h3>",
    ]
    if flagged:
        parts.append(_table(["Time", "Sender", "Recipient", "Subject", "Status", "Flag"],
                            [[f"{r['time']:%H:%M}", r["sender"], r["recipient"], r["subject"][:90],
                              r["status"], r["flags"]] for r in flagged[:60]]))
        if len(flagged) > 60:
            parts.append(f"<p>…and {len(flagged) - 60} more on the Flags sheet.</p>")
    else:
        parts.append("<p>Nothing flagged.</p>")
    parts += [
        "<h3>Busiest mailboxes</h3>",
        _table(["Mailbox", "Sent", "Ext. recipients", "Recv. internal", "Recv. external"],
               [[b["mailbox"], b["sent"], b["ext_rcpts"], b["recv_int"], b["recv_ext"]] for b in boxes[:20]]),
        "<p style=\"color:#555;font-size:12px\">Every message is listed in the attached workbook. "
        "Built from Exchange Online message trace metadata only; no message content was read.</p></div>",
    ]
    return "".join(parts)


def window():
    end_env = os.environ.get("WINDOW_END", "").strip()
    now = dt.datetime.now(BAGHDAD)
    if end_env == "now":
        end = now.replace(second=0, microsecond=0)
    elif end_env:
        end = dt.datetime.fromisoformat(end_env.replace("Z", "+00:00"))
        end = end if end.tzinfo else end.replace(tzinfo=BAGHDAD)
    else:
        end = now.replace(hour=20, minute=0, second=0, microsecond=0)
        if end > now:
            end -= dt.timedelta(days=1)
    return end - dt.timedelta(days=1), end.astimezone(BAGHDAD)


def main():
    start, end = window()
    start = start.astimezone(BAGHDAD)
    token = graph_token()
    rows = [classify(t) for t in message_traces(token, start, end)]
    rows = [r for r in rows if start <= r["time"] < end]
    rows.sort(key=lambda r: r["time"])
    boxes = mailbox_stats(rows)
    by_dir = collections.Counter(r["direction"] for r in rows)
    kpi = {
        "Message records (one per recipient)": len(rows),
        "Unique messages": len({r["message_id"] for r in rows}),
        "Internal": by_dir["Internal"],
        "Outbound (sent to external)": by_dir["Outbound"],
        "Inbound (received from external)": by_dir["Inbound"],
        "Not delivered": sum(r["status"] in ("Failed", "Quarantined", "Filtered as spam") for r in rows),
        "Flagged": sum(bool(r["flags"]) for r in rows),
        "Active mailboxes": len(boxes),
    }
    xlsx = build_workbook(rows, boxes, start, end, kpi)
    if len(xlsx) > MAX_ATTACHMENT:
        raise SystemExit(f"Workbook is {len(xlsx)} bytes, over the sendMail attachment limit")
    message = {
        "subject": f"Daily mail-flow report – {end:%d %b %Y} ({len(rows)} records, "
                   f"{kpi['Flagged']} flagged)",
        "body": {"contentType": "HTML", "content": build_html(rows, boxes, start, end, kpi)},
        "toRecipients": [{"emailAddress": {"address": a}} for a in RECIPIENTS],
        "attachments": [{
            "@odata.type": "#microsoft.graph.fileAttachment",
            "name": f"Mail-flow_{end:%Y-%m-%d}.xlsx",
            "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "contentBytes": base64.b64encode(xlsx).decode(),
        }],
    }
    graph(token, f"{GRAPH}/v1.0/users/{urllib.parse.quote(SENDER)}/sendMail",
          {"message": message, "saveToSentItems": False})
    # The repo is public, so the Actions log carries no mail data, not even counts.
    print(f"Report sent for the window ending {end:%Y-%m-%d %H:%M}")


if __name__ == "__main__":
    main()
